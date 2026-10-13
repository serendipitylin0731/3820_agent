from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from lxml import html

from config import Settings
from knowledge import knowledge_base as kb


def test_ppt_keeps_actual_slide_number(tmp_path):
    path = tmp_path / 'slides.ppt'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = '三年目标'
    prs.save(path)
    chunks = kb.load_ppt_chunks(path)
    assert len(chunks) == 1
    assert chunks[0].source.page_or_slide == 2
    assert chunks[0].content == '三年目标'


def test_word_uses_paragraph_and_table_positions(tmp_path):
    path = tmp_path / 'source.doc'
    document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>福州</w:t></w:r></w:p>'
                '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>指标</w:t></w:r></w:p></w:tc>'
                '<w:tc><w:p><w:r><w:t>1995</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
                '</w:body></w:document>')
    with ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', document)
    chunks = kb.load_word_chunks(path)
    assert [hit.content for hit in chunks] == ['福州', '指标 | 1995']
    assert [hit.source.page_or_slide for hit in chunks] == ['段落1', '段落2']


def test_media_and_ppt_sidecars_are_used_by_chat_and_page_ranking(tmp_path, monkeypatch):
    import json
    root = tmp_path
    docs = root / 'docs'
    docs.mkdir()
    guide = docs / 'guide.md'
    guide.write_text('# 1 Guide\n', encoding='utf-8')
    library = root / 'knowledge' / 'library'
    derived = library / 'derived'
    derived.mkdir(parents=True)
    entries = []
    for identifier, title, suffix, category, text in [
        ('img40', '原件扫描第40页', '.png', 'media', '# 原件扫描第40页\n\n闽江口规划'),
        ('img41', '原件扫描第41页', '.png', 'media', '# 原件扫描第41页\n\n闽江口规划'),
        ('video', '3820纪录片', '.mp4', 'media', '# 3820纪录片\n\n[音频 00:00:03] 闽江口规划'),
        ('slide', '3820视频', '.pptx', 'presentations', '# 3820视频\n\n## 第 3 页\n\n闽江口规划'),
    ]:
        path = derived / f'{identifier}.txt'
        path.write_text(text, encoding='utf-8')
        entries.append({'id': identifier, 'title': title, 'category': category,
                        'path': f'knowledge/library/{identifier}{suffix}',
                        'text_path': f'knowledge/library/derived/{identifier}.txt',
                        'searchable': category == 'media'})
    catalog = library / 'catalog.json'
    catalog.write_text(json.dumps({'entries': entries}), encoding='utf-8')
    monkeypatch.setattr(kb, 'GUIDE', guide)
    monkeypatch.setattr(kb, 'DATA_DIR', root / 'knowledge' / 'data')
    monkeypatch.setattr(kb, 'LIBRARY_CATALOG', catalog)
    kb._load_knowledge_base.cache_clear()
    try:
        paths = kb._source_paths(Settings())
        assert all((derived / f'{identifier}.txt').resolve() in paths
                   for identifier in ('img40', 'img41', 'video', 'slide'))
        chunks = kb.get_knowledge_base()
        assert kb.search('原件扫描第40页图片写了什么', chunks=chunks).hits[0].source.page_or_slide == 40
        assert kb.search('3820视频PPT第3页讲了什么', chunks=chunks).hits[0].source.section == '演示文稿'
        assert kb.search('3820纪录片视频讲了什么', chunks=chunks).hits[0].source.section == '视频语音与画面文字'
    finally:
        kb._load_knowledge_base.cache_clear()


def test_scan_ocr_page_retains_preview_page_and_unverified_status(tmp_path):
    path = tmp_path / 'page_0037.txt'
    path.write_text('1995 年主要指标翻一番', encoding='utf-8')
    chunks = kb.load_scan_page_chunk(path, {'article_url': 'https://example.test/original'})
    assert chunks[0].source.page_or_slide == '预览第37页'
    assert chunks[0].source.source_file == 'https://example.test/original'
    assert chunks[0].source.verified is False


def test_ppt_image_ocr_requires_matching_source_and_keeps_slide_number(tmp_path):
    import hashlib
    path = tmp_path / 'slides.pptx'
    path.write_bytes(b'present-version')
    text_path = tmp_path / 'abc' / 'slide_0019.txt'
    text_path.parent.mkdir()
    text_path.write_text('闽江口金三角经济圈投资示意图\n', encoding='utf-8')
    metadata = {'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'slides': [{'page': 19, 'text_file': 'abc/slide_0019.txt',
                            'text_sha256': hashlib.sha256(text_path.read_bytes()).hexdigest()}]}
    chunks = kb.load_ppt_image_ocr_chunks(path, metadata, tmp_path)
    assert len(chunks) == 1
    assert chunks[0].source.page_or_slide == 19
    assert chunks[0].source.section == '幻灯片图片 OCR'
    assert chunks[0].source.verified is False
    path.write_bytes(b'new-version')
    assert kb.load_ppt_image_ocr_chunks(path, metadata, tmp_path) == []


def test_official_yearbook_extracts_only_fuzhou_row():
    from knowledge.fetch_public_sources import _fuzhou_2010_yearbook_row
    page = html.fromstring('<table><tr><td>福州市</td><td>3123.41</td><td>282.73</td>'
                           '<td>1401.92</td><td>1127.59</td><td>274.33</td><td>1438.76</td></tr>'
                           '<tr><td>福州市辖区</td><td>1545.23</td></tr></table>')
    text = _fuzhou_2010_yearbook_row(page)
    assert '地区生产总值 3123.41' in text
    assert '第三产业增加值 1438.76' in text
    assert '1545.23' not in text


def test_public_source_decoder_respects_declared_charset():
    from knowledge.fetch_public_sources import _decode_page
    for charset, codec in [('utf-8', 'utf-8'), ('gb2312', 'gb18030')]:
        payload = (f'<html><head><meta charset="{charset}"></head><body>福州市</body></html>'
                   .encode(codec))
        assert '福州市' in ''.join(_decode_page(payload).itertext())


def test_vision_embedding_uses_multimodal_endpoint_and_query_instruction():
    class Client:
        def __init__(self):
            self.calls = []

        def post(self, path, **kwargs):
            self.calls.append((path, kwargs))
            return {'data': {'embedding': [1.0, 2.0]}}

    client = Client()
    settings = Settings(embed_model='doubao-embedding-vision-251215')
    assert kb.embed_texts(['文档'], settings=settings, client=client) == [[1.0, 2.0]]
    assert kb.embed_texts(['问题'], settings=settings, client=client, role='query') == [[1.0, 2.0]]
    assert all(path == '/embeddings/multimodal' for path, _ in client.calls)
    assert client.calls[0][1]['body']['instructions'] != client.calls[1][1]['body']['instructions']
