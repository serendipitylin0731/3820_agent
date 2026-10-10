"""对本地 PPT 图片形状 OCR，按实际幻灯片页号保存可复核文本。"""

import argparse
import hashlib
import json
from datetime import date
from io import BytesIO
from pathlib import Path

import numpy as np
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from rapidocr_onnxruntime import RapidOCR

DATA = Path(__file__).resolve().parent / 'data'
OUTPUT = DATA / 'ppt_ocr'


def _ppt_files(data):
    files = sorted(path for path in data.iterdir() if path.is_file() and
                   path.suffix.lower() in {'.ppt', '.pptx'})
    if any(path.name == '3820视频.pptx' for path in files):
        files = [path for path in files if path.name != '3820PPT+视频.ppt']
    return files


def _native_text(slide):
    parts = []
    for shape in slide.shapes:
        if getattr(shape, 'has_text_frame', False):
            parts.append(shape.text_frame.text)
    return ''.join(parts).replace(' ', '').replace('\n', '')


def _ocr_image(blob, engine):
    with Image.open(BytesIO(blob)) as image:
        image = image.convert('RGB')
        image.thumbnail((2400, 2400))
        if min(image.size) < 80:
            return []
        found, _ = engine(np.asarray(image))
    return [text.strip() for _box, text, confidence in found or []
            if confidence >= 0.55 and len(text.strip()) >= 2]


def extract(data=DATA, output=OUTPUT):
    data, output = Path(data), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    engine = RapidOCR()
    cache = {}
    decks = []
    for path in _ppt_files(data):
        deck_id = hashlib.sha256(path.name.encode('utf-8')).hexdigest()[:12]
        directory = output / deck_id
        directory.mkdir(parents=True, exist_ok=True)
        slides = []
        images = skipped = 0
        for page, slide in enumerate(Presentation(path).slides, 1):
            native = _native_text(slide)
            lines = []
            for shape in slide.shapes:
                if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    continue
                images += 1
                blob = shape.image.blob
                digest = hashlib.sha256(blob).hexdigest()
                if digest not in cache:
                    try:
                        cache[digest] = _ocr_image(blob, engine)
                    except (OSError, UnidentifiedImageError, ValueError):
                        cache[digest] = []
                        skipped += 1
                lines.extend(text for text in cache[digest]
                             if text.replace(' ', '') not in native)
            lines = list(dict.fromkeys(lines))
            content = '\n'.join(lines).strip()
            if len(content.replace('\n', '')) < 12:
                continue
            relative = f'{deck_id}/slide_{page:04d}.txt'
            (output / relative).write_text(content + '\n', encoding='utf-8')
            slides.append({'page': page, 'text_file': relative,
                           'text_sha256': hashlib.sha256((content + '\n').encode('utf-8')).hexdigest(),
                           'characters': len(content)})
        decks.append({'source_file': path.name,
                      'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                      'image_shapes': images, 'unreadable_images': skipped, 'slides': slides})
        print(f'{path.name}: {images} 张图片，{len(slides)} 页有新增 OCR 文本', flush=True)
    manifest = {'retrieved': date.today().isoformat(), 'ocr_engine': 'rapidocr-onnxruntime 1.4.4',
                'decks': decks}
    (output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return manifest


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data', type=Path, default=DATA)
    parser.add_argument('--output', type=Path, default=OUTPUT)
    options = parser.parse_args()
    result = extract(options.data, options.output)
    print(f"已处理 {len(result['decks'])} 份 PPT，共 "
          f"{sum(len(deck['slides']) for deck in result['decks'])} 页新增文字")
