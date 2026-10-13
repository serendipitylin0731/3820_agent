"""带来源的知识切片、离线文本检索和可选向量索引。"""
import hashlib
import json
import logging
import os
import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
from pathlib import Path
from threading import Lock
from xml.etree import ElementTree
from zipfile import ZipFile

from contracts import KnowledgeHit, KnowledgeResult, SourceRef
from config import Settings, get_settings
from resilience import transient_error
from tenacity import Retrying, retry_if_exception, stop_after_attempt, wait_exponential

GUIDE = Path(__file__).resolve().parent.parent / 'docs' / '基于3820战略工程内容的多智能体系统搭建指南.md'
DATA_DIR = Path(__file__).resolve().parent / 'data'
LIBRARY_CATALOG = Path(__file__).resolve().parent / 'library' / 'catalog.json'
WORD_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
EMBEDDING_SCHEMA = 2
ORIGINAL_ARTICLE_URL = 'https://article.xuexi.cn/articles/pdf/index.html?art_id=1854475750850039511'
EXACT_ENTITIES = ('长乐国际机场', '海上福州', '元洪投资区', '福州经济技术开发区')
_retrieval_counts = Counter()
_retrieval_lock = Lock()
logger = logging.getLogger('3820_agent')


def retrieval_metrics():
    with _retrieval_lock:
        return dict(_retrieval_counts)


def _record_retrieval(status, low_relevance=False):
    with _retrieval_lock:
        _retrieval_counts['requests'] += 1
        _retrieval_counts[status] += 1
        if low_relevance:
            _retrieval_counts['low_relevance'] += 1


def _version(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()[:16]


def load_markdown(path=GUIDE):
    path = Path(path)
    lines = path.read_text(encoding='utf-8').splitlines()
    chunks, heading, block, start = [], '', [], 1
    in_code = False
    started = False
    table_mode = False
    def flush():
        if not block:
            return
        content = '\n'.join(block).strip()
        if content:
            display_path = str(path.relative_to(GUIDE.parent.parent)) if path.is_relative_to(GUIDE.parent.parent) else path.name
            source = SourceRef(source_file=display_path, page_or_slide=f'L{start}',
                               section=heading or None, data_version=_version(path), verified=True)
            chunks.append(KnowledgeHit(chunk_id=f'{path.name}:{start}', content=content, source=source))
        block.clear()
    for line_no, line in enumerate(lines, 1):
        if not started:
            if re.match(r'^# 1\b', line):
                started = True
            else:
                continue
        if line.lstrip().startswith(('```', '~~~')):
            in_code = not in_code
            continue
        if line.startswith('<a id='):
            continue
        if in_code or line.startswith('# ') and re.match(r'# (?:3|4|5|6|7|8|9|10)\b', line):
            if not in_code and line.startswith('# '):
                break
            continue
        is_table_row = line.lstrip().startswith('|')
        if is_table_row and not table_mode:
            flush(); table_mode = True
        elif table_mode and not is_table_row:
            flush(); table_mode = False
        if re.match(r'^#{1,4} \d', line):
            flush(); heading = line.lstrip('# ').strip(); start = line_no
        if not block:
            start = line_no
        if line.strip():
            block.append(line.strip())
        if not table_mode and sum(map(len, block)) >= 700:
            flush()
    flush()
    return chunks


def load_ppt_text(path):
    from pptx import Presentation
    prs = Presentation(path)
    chunks = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, 'has_table', False):
                for row in shape.table.rows:
                    texts.append(' | '.join(cell.text.strip() for cell in row.cells))
        if texts:
            chunks.append(f'[第{idx}页] ' + '\n'.join(texts))
    return chunks


def load_ppt_chunks(path):
    path = Path(path)
    from pptx import Presentation
    version = _version(path)
    chunks = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, 'has_table', False):
                texts.extend(' | '.join(cell.text.strip() for cell in row.cells)
                             for row in shape.table.rows)
        content = '\n'.join(texts).strip()
        if content:
            chunks.append(KnowledgeHit(chunk_id=f'{path.name}:slide:{n}', content=content,
                                       source=SourceRef(source_file=_display_path(path), page_or_slide=n,
                                                        data_version=version, verified=False)))
    return chunks


def _display_path(path):
    path = Path(path)
    root = GUIDE.parent.parent
    return str(path.relative_to(root)) if path.is_relative_to(root) else path.name


def load_word_chunks(path):
    """读取 OOXML Word 的段落和表格；按文件内容识别，不依赖 .docx 后缀。"""
    path = Path(path)
    with ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read('word/document.xml'))
    body = root.find(f'{WORD_NS}body')
    if body is None:
        return []
    version = _version(path)
    chunks = []
    paragraph_no = 0
    for element in body:
        if element.tag == f'{WORD_NS}p':
            paragraph_no += 1
            content = ''.join(node.text or '' for node in element.iter(f'{WORD_NS}t')).strip()
        elif element.tag == f'{WORD_NS}tbl':
            rows = []
            for row in element.iter(f'{WORD_NS}tr'):
                cells = [''.join(t.text or '' for t in cell.iter(f'{WORD_NS}t')).strip()
                         for cell in row.iter(f'{WORD_NS}tc')]
                if any(cells):
                    rows.append(' | '.join(cells))
            content = '\n'.join(rows).strip()
            paragraph_no += 1
        else:
            continue
        if content:
            chunks.append(KnowledgeHit(chunk_id=f'{path.name}:paragraph:{paragraph_no}', content=content,
                                       source=SourceRef(source_file=_display_path(path),
                                                        page_or_slide=f'段落{paragraph_no}',
                                                        data_version=version, verified=False)))
    return chunks


def load_public_text_chunks(path, metadata):
    path = Path(path)
    content = path.read_text(encoding='utf-8')
    version = _version(path)
    chunks = []
    for n, paragraph in enumerate(content.split('\n\n'), 1):
        paragraph = paragraph.strip()
        if paragraph:
            chunks.append(KnowledgeHit(chunk_id=f'{path.stem}:paragraph:{n}', content=paragraph,
                                       source=SourceRef(source_file=metadata.get('url', _display_path(path)),
                                                        page_or_slide=f'段落{n}', section=metadata.get('title'),
                                                        data_version=version, verified=False)))
    return chunks


def load_scan_page_chunk(path, metadata):
    """OCR 页与扫描预览的实际页码绑定；识别文本不能自动视为核验事实。"""
    path = Path(path)
    content = path.read_text(encoding='utf-8').strip()
    if not content:
        return []
    page = int(path.stem.split('_')[-1])
    source = SourceRef(source_file=metadata['article_url'], page_or_slide=f'预览第{page}页',
                       section='原始扫描件 OCR', data_version=_version(path), verified=False)
    return [KnowledgeHit(chunk_id=f'original_3820:page:{page}', content=content, source=source)]


def load_ppt_image_ocr_chunks(path, metadata, root):
    """只读取与现有 PPT 校验值一致的 OCR，避免源文件变更后引用过期文本。"""
    if not metadata or metadata.get('source_sha256', '')[:16] != _version(path):
        return []
    chunks = []
    for slide in metadata['slides']:
        text_path = Path(root) / slide['text_file']
        if not text_path.is_file():
            continue
        content = text_path.read_text(encoding='utf-8').strip()
        if not content or hashlib.sha256(text_path.read_bytes()).hexdigest() != slide['text_sha256']:
            continue
        chunks.append(KnowledgeHit(chunk_id=f'{Path(path).name}:image-ocr:slide:{slide["page"]}',
                                   content=content,
                                   source=SourceRef(source_file=_display_path(path),
                                                    page_or_slide=slide['page'],
                                                    section='幻灯片图片 OCR',
                                                    data_version=_version(text_path), verified=False)))
    return chunks


def _source_paths(settings):
    # Teaching documents stay private; media and slides use prepared text sidecars.
    files = []
    indexed_scan_pages = set()
    extra = os.environ.get('AGENT3820_EXTRA_DATA_DIR')
    if extra and Path(extra).is_dir():
        files.extend(path for path in Path(extra).iterdir() if path.is_file() and
                     path.suffix.lower() in {'.txt', '.md'} and not
                     any(word in path.name for word in ('教学', '教案', '课件', '课程', '研修')))
    if LIBRARY_CATALOG.is_file():
        for item in json.loads(LIBRARY_CATALOG.read_text(encoding='utf-8')).get('entries', []):
            if item.get('category') in {'media', 'presentations'} and item.get('text_path'):
                path = GUIDE.parent.parent / item['text_path']
                if path.is_file():
                    files.append(path)
                    source = item.get('source', '')
                    if source.startswith('knowledge/data/original_3820/pages/'):
                        indexed_scan_pages.add((GUIDE.parent.parent / source).with_suffix('.txt').resolve())
            if item.get('searchable') and not item.get('source', '').startswith('knowledge/data/public_sources/') \
                    and item.get('category') in {'documents', 'papers', 'literature'} \
                    and item['path'].endswith('.md'):
                path = GUIDE.parent.parent / item['path']
                if path.is_file() and path != GUIDE:
                    files.append(path)
    public_dir = DATA_DIR / 'public_sources'
    if public_dir.is_dir():
        files.extend(public_dir.glob('*.txt'))
        if (public_dir / 'manifest.json').is_file():
            files.append(public_dir / 'manifest.json')
    scan_dir = DATA_DIR / 'original_3820'
    if scan_dir.is_dir():
        files.extend(path for path in (scan_dir / 'pages').glob('page_*.txt')
                     if path.resolve() not in indexed_scan_pages)
        for name in ('manifest.json', 'catalog.json'):
            if (scan_dir / name).is_file():
                files.append(scan_dir / name)
    files.append(LIBRARY_CATALOG) if LIBRARY_CATALOG.is_file() else None
    return tuple(sorted({path.resolve() for path in files}, key=str))


def get_knowledge_base():
    settings = get_settings()
    paths = _source_paths(settings)
    signatures = tuple((str(path), path.stat().st_size, path.stat().st_mtime_ns) for path in paths)
    return _load_knowledge_base(_version(GUIDE), signatures)


@lru_cache(maxsize=4)
def _load_knowledge_base(_guide_version, signatures):
    hits = load_markdown(GUIDE)
    media_by_text = {}
    if LIBRARY_CATALOG.is_file():
        media_by_text = {(GUIDE.parent.parent / item['text_path']).resolve(): item
                         for item in json.loads(LIBRARY_CATALOG.read_text(encoding='utf-8')).get('entries', [])
                         if item.get('category') in {'media', 'presentations'} and item.get('text_path')}
    public_dir = DATA_DIR / 'public_sources'
    manifest_path = public_dir / 'manifest.json'
    metadata = ({item['id']: item for item in json.loads(manifest_path.read_text(encoding='utf-8'))}
                if manifest_path.is_file() else {})
    scan_dir = DATA_DIR / 'original_3820'
    scan_manifest = scan_dir / 'manifest.json'
    scan_metadata = (json.loads(scan_manifest.read_text(encoding='utf-8'))
                     if scan_manifest.is_file() else None)
    for name, _, _ in signatures:
        path = Path(name)
        if path in media_by_text:
            hits.extend(load_derived_chunks(path, media_by_text[path]))
        elif path.suffix.lower() == '.md':
            hits.extend(load_library_markdown_chunks(path))
        elif path.suffix.lower() == '.txt' and path.parent == scan_dir / 'pages':
            if scan_metadata:
                hits.extend(load_scan_page_chunk(path, scan_metadata))
        elif path.suffix.lower() == '.txt':
            hits.extend(load_public_text_chunks(path, metadata.get(path.stem, {})))
    return hits


def load_derived_chunks(path, item):
    """Keep original media/PPT as the cited source, with slide or timestamp."""
    text = Path(path).read_text(encoding='utf-8-sig', errors='replace').strip()
    source_file = item['path']
    version = _version(path)
    suffix = Path(source_file).suffix.lower()
    chunks = []

    def add(content, position, section):
        if content.strip():
            chunks.append(KnowledgeHit(chunk_id=f'{item["id"]}:{len(chunks) + 1}',
                                       content=f'{item["title"]}\n{content.strip()}',
                                       source=SourceRef(source_file=source_file, page_or_slide=position,
                                                        section=section, data_version=version,
                                                        verified=False)))

    if suffix in {'.ppt', '.pptx'}:
        markers = list(re.finditer(r'(?m)^## 第 (\d+) 页\s*$', text))
        for index, marker in enumerate(markers):
            end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
            add(text[marker.end():end], int(marker.group(1)), '演示文稿')
    elif suffix in {'.mp4', '.mov', '.m4v'}:
        block, start = [], None
        for line in text.splitlines():
            match = re.match(r'^\[(音频|画面) (\d\d:\d\d:\d\d)\] (.+)$', line)
            if not match:
                continue
            if block and sum(map(len, block)) + len(line) > 700:
                add('\n'.join(block), start, '视频语音与画面文字')
                block = []
            if not block:
                start = match.group(2)
            block.append(line)
        if block:
            add('\n'.join(block), start, '视频语音与画面文字')
        if not chunks:
            add('未提取到可检索的语音或画面文字；可打开原视频查看。', None, '视频语音与画面文字')
    else:
        page = re.search(r'原件扫描第(\d+)页', item['title'])
        add(text.removeprefix(f"# {item['title']}").strip() or '未识别到文字；可打开原图查看。',
            int(page.group(1)) if page else None, '图片文字识别')
    return chunks


def load_library_markdown_chunks(path):
    """Split long Markdown transcriptions into bounded searchable passages."""
    content = Path(path).read_text(encoding='utf-8-sig', errors='replace')
    page_headers = list(re.finditer(r'(?m)^## 第 (\d+) 页(?:\s*·[^\n]*)?\s*$', content))
    hits, block, offset = [], '', 1
    version = _version(path)

    def append(text, number, page=None):
        if text.strip():
            hits.append(KnowledgeHit(chunk_id=f'{Path(path).name}:block:{number}', content=text.strip(),
                                     source=SourceRef(source_file=_display_path(path),
                                                      page_or_slide=(f'第 {page} 页' if page else f'片段{number}'),
                                                      data_version=version, verified=False)))

    def add_sections(text, page=None):
        nonlocal block, offset
        text = re.sub(r'(?m)^[ \t]*---[ \t]*$', '', text)
        for section in re.split(r'\n\s*\n', text):
            for piece in (section[i:i + 1200] for i in range(0, len(section), 1200)):
                if len(block) + len(piece) > 1200 and block:
                    append(block, offset, page)
                    offset += 1
                    block = ''
                block += ('\n\n' if block else '') + piece
        append(block, offset, page)
        offset += bool(block.strip())
        block = ''

    if page_headers:
        # Never blend adjacent pages: the chat citation must open the PDF on
        # the page from which the answer was retrieved.
        add_sections(content[:page_headers[0].start()])
        for index, header in enumerate(page_headers):
            end = page_headers[index + 1].start() if index + 1 < len(page_headers) else len(content)
            add_sections(content[header.end():end], int(header.group(1)))
    else:
        add_sections(content)
    return hits


def _tokens(text):
    chars = re.findall(r'[\u4e00-\u9fff]|[A-Za-z0-9]+', text.lower())
    return set(chars + [chars[i] + chars[i+1] for i in range(len(chars)-1)])


def _named_document(query):
    """Resolve a book named by the user before ranking its individual passages."""
    compact = re.sub(r'[\s《》“”·（）()]', '', query)
    if '生态文明思想概论' in compact:
        return '习近平生态文明思想概论.md'
    if '福建足迹' in compact or '闽山闽水物华新' in compact:
        if re.search(r'(?:下册|下卷|足迹下|物华新下)', compact):
            return '闽山闽水物华新·习近平福建足迹（下）.md'
        if re.search(r'(?:上册|上卷|足迹上|物华新上)', compact):
            return '闽山闽水物华新·习近平福建足迹（上）.md'
    if '谈治国理政' in compact:
        for volume in ('第一', '第二', '第三', '第四', '第五'):
            if volume + '卷' in compact:
                return f'习近平谈治国理政·{volume}卷.md'
    if '六中全会公报' in compact or '第十九届中央委员会第六次全体会议公报' in compact:
        return '中国共产党第十九届中央委员会第六次全体会议公报.md'
    if '十九大报告' in compact or '第十九次全国代表大会报告' in compact:
        return '中国共产党第十九次全国代表大会报告.md'
    return None


def _source_page(hit):
    position = hit.source.page_or_slide
    if isinstance(position, int):
        return position
    if isinstance(position, str):
        match = re.search(r'第\s*(\d+)\s*页', position)
        if match:
            return int(match.group(1))
    return None


def _exact_entity_hits(query, chunks):
    """用明确的实体名称补回向量检索可能漏掉的原文段落。"""
    anchors = [name for name in EXACT_ENTITIES if name in query]
    if not anchors:
        return []
    query_tokens = _tokens(query)
    found = []
    for anchor in anchors:
        candidates = [_expand_interview_hit(hit, chunks) for hit in chunks if anchor in hit.content]
        if not candidates:
            continue
        heading = re.compile(rf'(?m)^\s*(?:福州)?{re.escape(anchor)}\s*$')
        best = max(candidates, key=lambda hit: (
            bool(heading.search(hit.content)),
            hit.source.source_file == ORIGINAL_ARTICLE_URL,
            len(query_tokens & _tokens(hit.content)) / max(1, len(hit.content) ** 0.5),
            -hit.content.find(anchor),
        ))
        if best.chunk_id not in {item.chunk_id for item in found}:
            found.append(best.model_copy(update={'score': 1.0, 'score_kind': 'similarity'}))
    return found


def _expand_interview_hit(hit, chunks):
    """访谈问句和答复常被切成相邻段落，返回时合并供规划分析。"""
    marker = hit.source.page_or_slide
    if not isinstance(marker, str) or not marker.startswith('段落') or len(hit.content) > 130:
        return hit
    try:
        number = int(marker.removeprefix('段落'))
    except ValueError:
        return hit
    index = {other.source.page_or_slide: other for other in chunks
             if other.source.source_file == hit.source.source_file}
    next_hits = []
    for current in range(number + 1, number + 9):
        other = index.get(f'段落{current}')
        if other is None:
            break
        if next_hits and len(other.content) < 130 and any(mark in other.content for mark in ('？', '?')):
            break
        next_hits.append(other)
    if not next_hits:
        return hit
    content = '\n'.join([hit.content] + [other.content for other in next_hits])[:3600]
    source = hit.source.model_copy(update={'page_or_slide': f'段落{number}–{number+len(next_hits)}'})
    return hit.model_copy(update={'content': content, 'source': source})


def search(query, *, chunks=None, top_k=5, min_score=0.35):
    if top_k < 1 or not query.strip():
        _record_retrieval('insufficient_evidence')
        return KnowledgeResult(status='insufficient_evidence')
    settings = get_settings()
    if chunks is None and not LIBRARY_CATALOG.is_file() and settings.vector_max_distance is not None and (settings.index_path / 'index.faiss').is_file():
        try:
            vector = vector_search(query, top_k=top_k, max_distance=settings.vector_max_distance,
                                   directory=settings.index_path)
        except (ValueError, OSError) as exc:
            logger.warning('向量索引不可用，改用文本检索 error_type=%s', type(exc).__name__)
        else:
            exact = _exact_entity_hits(query, get_knowledge_base())
            merged = {hit.chunk_id: hit for hit in [*exact, *vector.hits]}
            expanded = [_expand_interview_hit(hit, get_knowledge_base()) for hit in list(merged.values())[:top_k]]
            result = KnowledgeResult(status='ok', hits=expanded) if expanded else \
                KnowledgeResult(status='insufficient_evidence')
            _record_retrieval(result.status)
            return result
    chunks = get_knowledge_base() if chunks is None else chunks
    target_name = _named_document(query)
    if target_name:
        named_chunks = [hit for hit in chunks if Path(hit.source.source_file).name == target_name]
        if named_chunks:
            chunks = named_chunks
            page_match = re.search(r'第\s*(\d+)\s*页', query)
            if page_match:
                requested = int(page_match.group(1))
                page_hits = [hit.model_copy(update={'score': 1.0, 'score_kind': 'similarity'})
                             for hit in chunks if _source_page(hit) == requested]
                result = KnowledgeResult(status='ok', hits=page_hits[:top_k]) if page_hits else \
                    KnowledgeResult(status='insufficient_evidence')
                _record_retrieval(result.status)
                return result
    query_tokens = _tokens(query)
    media_request = any(term in query for term in
                        ('示意图', '地图', '幻灯片', 'PPT', '课件', '图片', '图像', '照片', '扫描件', '视频', '录像', '画面'))
    # Explicit format and page references are more reliable than character overlap
    # with lengthy OCR/transcripts.  In particular, "第40页" must retrieve page 40.
    if re.search(r'PPT|ppt|幻灯片|演示文稿|课件', query):
        requested_section = '演示文稿'
    elif any(word in query for word in ('视频', '录像', '影片')):
        requested_section = '视频语音与画面文字'
    elif any(word in query for word in ('图片', '图像', '照片', '扫描件', '扫描页')):
        requested_section = '图片文字识别'
    else:
        requested_section = None
    page_match = re.search(r'第\s*(\d+)\s*页', query)
    requested_page = int(page_match.group(1)) if page_match else None
    ranked = []
    for hit in chunks:
        overlap = len(query_tokens & _tokens(hit.content))
        score = overlap / max(1, len(query_tokens))
        if media_request and hit.source.section in {
                '图片文字识别', '视频语音与画面文字', '演示文稿'}:
            if requested_section == hit.source.section:
                score += 0.9
            elif requested_section:
                score -= 0.7
            else:
                score += 0.2
            if requested_page is not None:
                score += 2.0 if _source_page(hit) == requested_page else -0.8
            title = hit.content.split('\n', 1)[0]
            deck = re.search(r'《([^》]+)》内嵌视频', title)
            name = deck.group(1) if deck else Path(hit.source.source_file).stem
            if len(name) >= 4 and name in query:
                score += 1.0
        if score >= min_score:
            ranked.append((score, hit))
    ranked.sort(key=lambda x: x[0], reverse=True)
    if requested_section and any(hit.source.section == requested_section for _, hit in ranked):
        ranked = [(score, hit) for score, hit in ranked if hit.source.section == requested_section]
    if requested_page is not None and any(_source_page(hit) == requested_page for _, hit in ranked):
        ranked = [(score, hit) for score, hit in ranked if _source_page(hit) == requested_page]
    if requested_section and ranked:
        def named_source(hit):
            title = hit.content.split('\n', 1)[0]
            match = re.search(r'《([^》]+)》内嵌视频', title)
            name = match.group(1) if match else Path(hit.source.source_file).stem
            return len(name) >= 4 and name in query
        if any(named_source(hit) for _, hit in ranked):
            ranked = [(score, hit) for score, hit in ranked if named_source(hit)]
    hits = [hit.model_copy(update={'score': round(score, 4), 'score_kind': 'similarity'})
            for score, hit in ranked[:top_k]]
    result = KnowledgeResult(status='ok', hits=hits) if hits else KnowledgeResult(status='insufficient_evidence')
    _record_retrieval(result.status, low_relevance=not hits and any(_tokens(query) & _tokens(hit.content) for hit in chunks))
    return result


def format_context(result):
    if result.status != 'ok':
        return '知识库暂无充分依据。'
    return '\n\n'.join(f'[{hit.source.section or "资料"} {hit.source.page_or_slide}] '
                       f'{hit.content}' for hit in result.hits)


def retrieve_context(query, top_k=None):
    result = search(query, top_k=top_k or get_settings().top_k)
    return format_context(result)


def get_client(settings: Settings | None = None):
    from openai import OpenAI
    settings = settings or get_settings()
    return OpenAI(base_url=settings.ark_base_url, api_key=settings.require_api_key(),
                  timeout=settings.request_timeout)


def embed_texts(texts, *, settings: Settings | None = None, client=None, role='document'):
    settings = settings or get_settings()
    active_client = client or get_client(settings)
    model = settings.require_embed_model()
    if model.startswith('doubao-embedding-vision-'):
        if role not in {'document', 'query'}:
            raise ValueError('向量角色必须是 document 或 query')
        instructions = ('Target_modality: text.\nInstruction:从资料中检索能回答问题的内容\nQuery:'
                        if role == 'query' else
                        'Target_modality: text.\nInstruction:为中文资料建立可供问答检索的索引\nQuery:')
        def embed_one(item):
            body = {'model': model, 'input': [{'type': 'text', 'text': item}],
                    'encoding_format': 'float'}
            if model >= 'doubao-embedding-vision-251215':
                body['instructions'] = instructions
            for attempt in Retrying(stop=stop_after_attempt(3), wait=wait_exponential(multiplier=1, max=8),
                                    retry=retry_if_exception(transient_error), reraise=True):
                with attempt:
                    response = active_client.post('/embeddings/multimodal', cast_to=dict, body=body,
                                                  options={'timeout': settings.request_timeout})
            embedding = response.get('data', {}).get('embedding')
            if not isinstance(embedding, list) or not embedding:
                raise ValueError('向量服务未返回稠密向量')
            return embedding
        if len(texts) < 2:
            return [embed_one(item) for item in texts]
        with ThreadPoolExecutor(max_workers=min(4, len(texts))) as executor:
            return list(executor.map(embed_one, texts))
    response = active_client.embeddings.create(model=model, input=texts,
                                               timeout=settings.request_timeout)
    return [item.embedding for item in response.data]


def build_index(chunks):
    import faiss
    import numpy as np
    if not chunks:
        raise ValueError('知识库没有可索引的文本片段')
    vectors = np.asarray(embed_texts(chunks), dtype='float32')
    if vectors.ndim != 2 or vectors.shape[0] != len(chunks) or not vectors.shape[1]:
        raise ValueError('向量数量或维度与知识片段不一致')
    index = faiss.IndexFlatL2(vectors.shape[1]); index.add(vectors)
    return index


def retrieve(index, chunks, query, top_k=5):
    import numpy as np
    if index is None or not chunks or not index.ntotal or top_k <= 0:
        return []
    top_k = min(top_k, index.ntotal, len(chunks))
    vector = np.asarray(embed_texts([query], role='query'), dtype='float32')
    if vector.ndim != 2 or vector.shape[1] != index.d:
        raise ValueError('查询向量维度与索引不一致')
    distances, ids = index.search(vector, top_k)
    return [(chunks[i], float(distances[0][j])) for j, i in enumerate(ids[0]) if 0 <= i < len(chunks)]


def rebuild_vector_store(directory=None):
    """基于当前资料生成向量索引和完整来源映射。"""
    settings = get_settings()
    hits = get_knowledge_base()
    index = build_index([hit.content for hit in hits])
    directory = directory or settings.index_path
    import faiss
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    faiss.write_index(index, str(directory / 'index.faiss'))
    metadata = {'model': settings.require_embed_model(), 'embedding_schema': EMBEDDING_SCHEMA,
                'hits': [hit.model_dump() for hit in hits]}
    (directory / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False), encoding='utf-8')
    return len(hits)


def vector_search(query, *, top_k=5, max_distance, directory=None):
    """索引版本、模型和片段位置必须与当前资料一致。"""
    import faiss
    import numpy as np
    settings = get_settings()
    directory = Path(directory or settings.index_path)
    metadata = json.loads((directory / 'metadata.json').read_text(encoding='utf-8'))
    current = get_knowledge_base()
    saved = [KnowledgeHit.model_validate(item) for item in metadata['hits']]
    if (metadata['model'] != settings.require_embed_model() or
            metadata.get('embedding_schema') != EMBEDDING_SCHEMA or
            [h.model_dump() for h in saved] != [h.model_dump() for h in current]):
        raise ValueError('资料或向量模型已变化，请重建索引')
    index = faiss.read_index(str(directory / 'index.faiss'))
    if index.ntotal != len(saved):
        raise ValueError('索引和来源映射数量不一致')
    vector = np.asarray(embed_texts([query], role='query'), dtype='float32')
    if vector.ndim != 2 or vector.shape[1] != index.d:
        raise ValueError('查询向量维度与索引不一致')
    distances, ids = index.search(vector, min(max(top_k * 6, top_k), index.ntotal))
    candidates = [saved[i].model_copy(update={'score': float(distances[0][j]), 'score_kind':'l2_distance'})
                  for j, i in enumerate(ids[0]) if i >= 0 and distances[0][j] <= max_distance]
    hits = _prefer_primary_source(candidates, top_k)
    return KnowledgeResult(status='ok', hits=hits) if hits else KnowledgeResult(status='insufficient_evidence')


def _prefer_primary_source(candidates, top_k):
    """当原件相关性接近首位时，至少保留一条原件；分数仍显示原始距离。"""
    if not candidates or top_k < 1:
        return []
    nearest = candidates[0].score
    primary = next((hit for hit in candidates
                    if hit.source.source_file == ORIGINAL_ARTICLE_URL and hit.score <= nearest + 0.2), None)
    if primary is None:
        return candidates[:top_k]
    return [primary] + [hit for hit in candidates if hit.chunk_id != primary.chunk_id][:top_k - 1]


if __name__ == '__main__':
    print(f'索引片段数：{rebuild_vector_store()}')
